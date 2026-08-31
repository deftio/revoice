"""Extract plain text from corpus files. One-way (corpus ingestion only, not output round-trip)."""

from __future__ import annotations

from pathlib import Path

TEXT_EXTS = {".md", ".txt", ".markdown", ".text", ".rst"}
SKIP_NAMES = {".DS_Store"}


def extract_text(path: Path) -> str | None:
    """Return document text, or None if unsupported."""
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        return path.read_text(errors="replace")
    if ext == ".docx":
        try:
            import docx  # type: ignore
        except ImportError:
            return None
        d = docx.Document(str(path))
        return "\n\n".join(p.text for p in d.paragraphs if p.text.strip())
    if ext == ".pptx":
        try:
            from pptx import Presentation  # type: ignore
        except ImportError:
            return None
        prs = Presentation(str(path))
        chunks = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    chunks.append(shape.text_frame.text)
        return "\n\n".join(chunks)
    if ext == ".pdf":
        try:
            import pypdfium2 as pdfium  # type: ignore
        except ImportError:
            return None
        pdf = pdfium.PdfDocument(str(path))
        return "\n\n".join(page.get_textpage().get_text_range() for page in pdf)
    return None


def walk_corpus(training_dir: Path):
    """Yield (path, relative_path_str) for every candidate file."""
    for p in sorted(training_dir.rglob("*")):
        if p.is_file() and p.name not in SKIP_NAMES and not p.name.startswith("."):
            yield p, str(p.relative_to(training_dir))
