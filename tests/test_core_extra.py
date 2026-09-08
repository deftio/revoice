"""Edge cases across core modules: ingest, metrics, indexer, report, segments,
spans, style, lint, voicepack, helptext, preflight, stylometry, rubric."""

import json
import sys

import pytest

from revoice.core import lint as lintmod
from revoice.core.ingest import extract_text, walk_corpus
from revoice.core.metrics import build_baselines, score_against_pack, score_text
from revoice.core.report import build_report, render_html, render_json, render_md
from revoice.core.segments import analyze_blend, classify_segment
from revoice.core.spans import parse, reassemble, rewritable
from revoice.core.voicepack import VoicePack, list_voices
from revoice.helptext import TOPICS, topic_list
from revoice.voicemetric.features import cosine, fingerprint
from tests.conftest import FakeProvider

MINIMAL_PDF = b"""%PDF-1.4
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj
4 0 obj<</Length 44>>stream
BT /F1 24 Tf 72 720 Td (Hello PDF text) Tj ET
endstream
endobj
5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
trailer<</Root 1 0 R>>
%%EOF"""


# ---------- ingest ----------

def test_extract_txt_and_unknown(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("hello")
    assert extract_text(f) == "hello"
    assert extract_text(tmp_path / "a.xyz") is None


def test_extract_docx(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("First para")
    d.add_paragraph("   ")
    d.add_paragraph("Second para")
    p = tmp_path / "a.docx"
    d.save(str(p))
    assert extract_text(p) == "First para\n\nSecond para"


def test_extract_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Slide text here"
    p = tmp_path / "a.pptx"
    prs.save(str(p))
    assert "Slide text here" in extract_text(p)


def test_extract_pdf(tmp_path):
    p = tmp_path / "a.pdf"
    p.write_bytes(MINIMAL_PDF)
    assert "Hello PDF text" in extract_text(p)


@pytest.mark.parametrize("mod,ext", [("docx", ".docx"), ("pptx", ".pptx"), ("pypdfium2", ".pdf")])
def test_extract_import_fallbacks(tmp_path, monkeypatch, mod, ext):
    monkeypatch.setitem(sys.modules, mod, None)  # forces ImportError
    f = tmp_path / f"a{ext}"
    f.write_bytes(b"whatever")
    assert extract_text(f) is None


def test_walk_corpus_skips(tmp_path):
    (tmp_path / "keep.md").write_text("x")
    (tmp_path / ".DS_Store").write_text("x")
    (tmp_path / ".hidden.md").write_text("x")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "b.txt").write_text("y")
    rels = [rel for _, rel in walk_corpus(tmp_path)]
    assert rels == ["keep.md", "sub/b.txt"]


# ---------- metrics ----------

def test_score_against_pack_errors(tmp_path, learned_pack):
    empty = VoicePack.create(tmp_path / "d2", "empty")
    with pytest.raises(RuntimeError, match="no baselines"):
        score_against_pack("text", empty)
    with pytest.raises(RuntimeError, match="not in voice"):
        score_against_pack("text", learned_pack, register="nope")
    r = score_against_pack("The impedance matching network transforms the load.", learned_pack)
    assert r["best_register"] in r["results"]
    r2 = score_against_pack("text", learned_pack, register=r["best_register"])
    assert set(r2["results"]) == {r["best_register"]}


def test_build_baselines_skips_registers_with_no_readable_text(tmp_path):
    """A baseline is built from TEXT, never from cached index fingerprints alone.

    An indexed file whose text can no longer be read (deleted, or an extension whose
    support went away) used to still yield a baseline from its stored fingerprint —
    one with empty n-gram and tf-idf centroids, which scores garbage against anything.
    Skipping the register is the honest outcome; `learn` then reports it as missing
    rather than shipping a measure that cannot work.
    """
    pack = VoicePack.create(tmp_path / "d", "odd")
    # indexed, but no readable file on disk -> no baseline at all
    pack.append_index({"path": "x.bin", "register": "nofp", "polish": "draft",
                       "domains": [], "chars": 5, "hash": "h1", "fingerprint": None})
    pack.append_index({"path": "y.bin", "register": "thin", "polish": "draft",
                       "domains": [], "chars": 5, "hash": "h2",
                       "fingerprint": {"function_word_freq": {}, "punct_per_sentence": {}}})
    # a real doc plus a phantom unreadable one in the same register: the register
    # survives on the readable document alone
    (pack.training_dir / "real.md").write_text("Words appear in this real document for testing baselines.")
    pack.append_index({"path": "real.md", "register": "mix", "polish": "polished",
                       "domains": [], "chars": 50, "hash": "h3",
                       "fingerprint": fingerprint("Words appear in this real document for testing baselines.")})
    pack.append_index({"path": "ghost.bin", "register": "mix", "polish": "polished",
                       "domains": [], "chars": 50, "hash": "h4",
                       "fingerprint": fingerprint("another doc entirely")})
    seen = []
    baselines = build_baselines(pack, progress=lambda r, m: seen.append(r))
    assert "nofp" not in baselines
    assert "thin" not in baselines
    # progress fires once while building the baseline and once while calibrating it
    assert "mix" in baselines and set(seen) == {"mix"}
    assert baselines["mix"]["doc_count"] == 1
    assert baselines["mix"]["char_ngrams"]  # centroids are populated, not empty


def test_baseline_from_texts_with_no_texts():
    from revoice.core.metrics import baseline_from_texts

    assert baseline_from_texts([]) == {}


def test_mean_std_empty():
    from revoice.core.metrics import _mean_std

    assert _mean_std([]) == (0.0, 0.0)


def test_build_profiles_branches(tmp_path):
    from revoice.core.profiles import build_profiles

    pack = VoicePack.create(tmp_path / "d", "v")
    (pack.training_dir / "real.md").write_text("Actual text for the exemplar bank right here.")
    pack.append_index({"path": "real.md", "register": "professional", "polish": "polished",
                       "domains": [], "chars": 40, "hash": "h1", "fingerprint": None})
    pack.append_index({"path": "ghost.bin", "register": "professional", "polish": "polished",
                       "domains": [], "chars": 40, "hash": "h2", "fingerprint": None})
    boom = FakeProvider(lambda s, u, n: (_ for _ in ()).throw(RuntimeError("no profile")))
    msgs = []
    built = build_profiles(pack, boom, progress=lambda r, m: msgs.append(m))
    assert built == {} and any("profile ERROR" in m for m in msgs)
    assert pack.manifest()["readiness"] == "not_ready"


def test_score_text_empty_baseline():
    """A baseline with nothing in it must score 0 everywhere, not divide by zero."""
    from revoice.core.metrics import COMPONENTS

    r = score_text("hi there", {"function_words": {}})
    assert r["burrows_delta"] == 99.0
    assert set(r["components"]) == set(COMPONENTS)
    for name, value in r["components"].items():
        assert value == 0.0, f"{name} should be 0 against an empty baseline"
    assert r["composite"] == 0.0


def test_cosine_empty():
    assert cosine({}, {"a": 1.0}) == 0.0
    assert cosine({"a": 1.0}, {"a": 1.0}) == pytest.approx(1.0)


# ---------- indexer ----------

def test_indexer_branches(tmp_path, stub):
    from revoice.core.indexer import index_corpus

    pack = VoicePack.create(tmp_path / "d", "v")
    (pack.training_dir / "good.md").write_text("A solid professional document about circuits.")
    (pack.training_dir / "empty.md").write_text("   ")
    (pack.training_dir / "blob.bin").write_bytes(b"\x00")
    lines = []
    stats = index_corpus(pack, stub, progress=lambda a, b: lines.append((a, b)))
    assert stats["classified"] == 1 and stats["skipped"] == 1 and stats["unsupported"] == 1
    # re-run: unchanged hash skipped
    stats2 = index_corpus(pack, stub)
    assert stats2["classified"] == 0 and stats2["skipped"] == 2

    # classifier error branch + new register discovery
    (pack.training_dir / "new.md").write_text("Fresh text to classify.")
    boom = FakeProvider(lambda s, u, n: (_ for _ in ()).throw(RuntimeError("nope")))
    errs = []
    stats3 = index_corpus(pack, boom, progress=lambda a, b: errs.append(b))
    assert stats3["errors"] == 1 and any("ERROR" in e for e in errs)

    weird = FakeProvider(lambda s, u, n: json.dumps(
        {"register": "Weird_New", "polish": "draft", "domains": ["X"], "summary": "s", "confidence": 1}))
    stats4 = index_corpus(pack, weird)
    assert stats4["classified"] == 1
    assert pack.read_index()["new.md"]["register"] == "weird_new"


# ---------- report ----------

def test_report_renderings(learned_pack):
    text = "Short document. It has two sentences."
    rep = build_report(text, source="x.md")
    md = render_md(rep)
    assert "Voice match" not in md
    html = render_html(rep)
    assert "Voice match" not in html
    assert json.loads(render_json(rep))["source"] == "x.md"

    rep2 = build_report(text, source="x.md", pack=learned_pack)
    md2 = render_md(rep2)
    assert "Voice match" in md2 and "scores are noisy" in md2
    html2 = render_html(rep2)
    assert "Voice match" in html2


def test_report_blended_html():
    text = ('"Hello there," she said, and he walked away into the night quietly.\n\n'
            "The taxation implementation demonstrates the administration of the arrangement "
            "and the tabulation of the information in the presentation of the documentation.\n\n"
            "- alpha beta gamma delta\n- epsilon zeta eta theta\n- iota kappa lambda mu\n- nu xi omicron pi")
    rep = build_report(text)
    assert "verdict" in render_html(rep)


# ---------- segments ----------

def test_classify_segment_types():
    assert classify_segment("12345 67890 13579 24680 987654321 1122334455 999") == "numeric"
    assert classify_segment("short bit") == "fragment"
    code = "def f(x):\n    return x\nimport os\nvoid main() { }\n-> :: =="
    assert classify_segment(code + " padding words to pass twelve word minimum easily now") == "code"
    lst = "- alpha beta gamma delta\n- epsilon zeta eta theta\n- iota kappa lambda mu"
    assert classify_segment(lst) == "list"
    nar = '"Hello," she said. He walked away and they saw her leave. She thought about him.'
    assert classify_segment(nar) == "narrative"
    exp = ("The implementation of the specification requires consideration of the "
           "documentation and the administration of the configuration management.")
    assert classify_segment(exp) == "expository"


def test_analyze_blend_edges():
    assert analyze_blend("")["blend"] == {}
    assert analyze_blend("")["verdict"] == "mostly-uniform"
    # leading fragment folds into following (else branch), later fragment folds back
    out = analyze_blend("Tiny head\n\n" + "A longer expository paragraph about the implementation "
                        "of the configuration and the documentation of the administration.\n\nstub tail")
    assert out["segments"][0]["type"] == "expository"
    # fenced code protected
    out2 = analyze_blend("```\ncode here\n```\n\nProse paragraph follows with plenty of "
                         "explanation about the implementation of the notation system.")
    assert out2["blend"].get("code")
    # heterogeneity computed across long prose windows
    a = ("Short punchy line. " * 40)
    b = ("The administration of the implementation requires extraordinarily comprehensive "
         "documentation considerations throughout every conceivable organizational arrangement. " * 20)
    out3 = analyze_blend(a + "\n\n" + b)
    assert out3["style_heterogeneity"] > 0


# ---------- spans ----------

def test_spans_edges():
    doc = (
        "# Title\n"
        "\n"
        "Paragraph line one\ncontinues here.\n"
        "---\n"
        "| a | b |\n"
        "<div>html</div>\n"
        "> quoted first\n"
        "> quoted second\n"
        ">\n"
        "- item one\n"
        "  continuation of item\n"
        "- item two\n"
        "1. numbered\n"
        "```\nfence never closes\n"
    )
    spans = parse(doc)
    kinds = [s.kind for s in spans]
    assert "heading" in kinds and "quote" in kinds and "listitem" in kinds
    q = [s for s in spans if s.kind == "quote"][0]
    assert q.text == "quoted first quoted second"
    li = [s for s in spans if s.kind == "listitem"][0]
    assert li.text == "item one continuation of item"
    out = reassemble(spans)
    assert "# Title" in out and "| a | b |" in out and "---" in out
    assert "> quoted first quoted second" in out
    assert "fence never closes" in out
    assert rewritable(spans)
    # empty text
    assert reassemble(parse("")) == ""


def test_spans_closed_fence_and_para_stop():
    doc = "```py\nx = 1\n```\nA paragraph\nthat stops here\n# Head\n"
    spans = parse(doc)
    assert spans[0].raw == "```py\nx = 1\n```"
    para = [s for s in spans if s.kind == "para"][0]
    assert para.text == "A paragraph that stops here"


# ---------- style ----------

def test_style_module(tmp_path):
    from revoice.core.style import apply_hard_swaps, load_style, render_prompt_section, write_template

    assert load_style(tmp_path) == {}
    assert render_prompt_section({}) == ""
    assert write_template(tmp_path) is True
    assert write_template(tmp_path) is False
    style = {
        "rules": ["Start with the point."],
        "phrases": {"In any event,": ["Moving on,"]},
        "lexicon": {"car": {"automobile": 0.8, "car": 0.16, "ride": 0.03, "whip": 0.01}},
        "swaps": {"utilize": "use"},
    }
    sec = render_prompt_section(style)
    assert "Voice rules:" in sec and 'prefers "automobile"' in sec
    assert 'often "car"' in sec and 'occasionally "ride"' in sec and 'rarely "whip"' in sec
    new, applied = apply_hard_swaps("Utilize this. We utilize that.", style)
    assert new == "Use this. We use that."
    assert applied == ["utilize->use"]
    assert apply_hard_swaps("nothing to do", style) == ("nothing to do", [])


def test_style_load_real_file(tmp_path):
    from revoice.core.style import load_style

    (tmp_path / "style.yaml").write_text("rules:\n  - one\nswaps:\n")
    assert load_style(tmp_path) == {"rules": ["one"]}


# ---------- lint ----------

def test_lint_score_and_custom(tmp_path):
    pats = lintmod.load_patterns()
    assert lintmod.score("We delve into and leverage the tapestry.", pats) > 0
    assert lintmod.score("", pats) == 0.0
    (tmp_path / "lint.yaml").write_text("patterns:\n  - 'my banned phrase'\n")
    pats2 = lintmod.load_patterns(tmp_path)
    assert any(n.startswith("custom:") for n in lintmod.scan("MY BANNED PHRASE here", pats2))


# ---------- voicepack ----------

def test_voicepack_misc(tmp_path):
    p = VoicePack(tmp_path / "novoice")
    assert not p.exists()
    m = p.manifest()  # missing manifest file
    assert m["readiness"] == "not_ready"
    p2 = VoicePack.create(tmp_path / "data", "v", metadata={"a": 1})
    assert p2.manifest()["metadata"] == {"a": 1}
    p2.update_manifest(readiness="prompts_ready")
    assert p2.manifest()["readiness"] == "prompts_ready"
    # index with blank lines
    p2.index_path.write_text('{"path": "a", "x": 1}\n\n{"path": "a", "x": 2}\n')
    assert p2.read_index()["a"]["x"] == 2
    assert list_voices(tmp_path / "missing") == []
    assert [v.name for v in list_voices(tmp_path / "data")] == ["v"]


# ---------- helptext ----------

def test_helptext():
    out = topic_list()
    for name in TOPICS:
        assert name in out


# ---------- preflight ----------

def _long(s, n=30):
    return " ".join(s for _ in range(n))


def test_preflight_doc_kinds():
    from revoice.core.preflight import build_plan

    code = "```\n" + "x = 1\n" * 30 + "```"
    assert build_plan(code)["doc_kind"] == "data-or-code-heavy"

    acad = ("The taxonomy of the classification follows Smith et al. (2019) and prior "
            "implementations [1] [2] [3]. " * 6)
    assert build_plan(acad)["doc_kind"] == "academic-paper-like"

    fic = ('"Where are we going," she said. He took her hand and they went into the dark. ' * 10)
    assert build_plan(fic)["doc_kind"] == "fiction-dialogue"

    nar = ("She walked to the harbor and he followed her past the boats. They saw the storm "
           "come in and thought about the crossing they had made. " * 6)
    assert build_plan(nar)["doc_kind"] == "narrative-prose"

    notes = "\n".join(f"- item {w} alpha beta gamma delta epsilon" for w in "abcdefghij")
    assert build_plan(notes)["doc_kind"] == "notes-or-outline"

    tech = ("The gain equals 3 dB at 10 kHz and the offset = 5 mV across the 2 ms window "
            "while the supply holds 3 V under the 100 mA load in every measured configuration "
            "of the amplifier documentation and implementation. " * 3)
    assert build_plan(tech)["doc_kind"] == "technical-quantitative"

    exp = ("The implementation of the specification requires consideration of the "
           "documentation and administration. " * 5)
    assert build_plan(exp)["doc_kind"] == "expository-prose"


def test_preflight_chunk_seams_and_segments(learned_pack):
    from revoice.core.preflight import _chunk_seams, build_plan

    para = ("A paragraph of expository prose about the implementation of the "
            "documentation and its administration across the organization. ")
    big = "\n\n".join(para * 4 for _ in range(20))
    seams = _chunk_seams(big)
    assert seams and all(isinstance(s, int) for s in seams)
    with_headings = "\n\n".join(f"# Section {i}\n\n" + para * 4 for i in range(20))
    assert _chunk_seams(with_headings)
    nogaps = "x" * 20000  # no headings, no paragraph breaks -> seam = cursor
    assert _chunk_seams(nogaps) == [6000, 12000, 18000]
    assert _chunk_seams("short") == []

    doc = (
        "See https://example.com and Smith et al. (2020) for details [1].\n\n"
        "The offset = 3 mV at 10 kHz with 5 dB of gain across 2 ms and 7 V rails "
        "= 9 measurements per 4 Hz bin in the 6 ms capture window today.\n\n"
        '"Stay close," she said. "The tide turns fast." He nodded and they walked on '
        "through the fog to the boats waiting in the dark harbor below the cliff.\n\n"
        "- tiny\n- list\n\n"
        "12345 67890 13579 24680 99887 77665 54433 22110 998877 665544\n\n"
        + (learned_pack.training_dir / "doc1.md").read_text().strip() + "\n\n"
        "   \n\n"
    )
    from revoice.core.metrics import load_baselines

    reg = sorted(load_baselines(learned_pack))[0]
    plan = build_plan(doc, learned_pack, reg)
    assert plan["target_register"] == reg
    treatments = {s["treatment"] for s in plan["segments"]}
    assert "protect" in treatments and "cleanup-only" in treatments
    assert "pass-through (on-target)" in treatments  # corpus text scores above floor
    cautions = [c for s in plan["segments"] for c in s.get("cautions", [])]
    assert any("URL" in c for c in cautions) and any("citation" in c for c in cautions)
    assert any("math" in c for c in cautions) and any("dialogue" in c for c in cautions)
    # plan without pack
    plan2 = build_plan(doc)
    assert "target_register" not in plan2


# ---------- rubric (generic engine) ----------

def test_rubric_engine_edges(tmp_path):
    from revoice.rubric import judge_all, judge_dimension, load_rubrics_file, vote_metrics

    f = tmp_path / "r.yaml"
    f.write_text(
        "dimensions:\n"
        "  tone:\n"
        "    description: tone check\n"
        "    choices: {good: fine, bad: poor}\n"
        "    examples: [an example]\n"
        "    counter_examples: [a counter]\n"
    )
    rubrics = load_rubrics_file(f)
    assert "tone" in rubrics

    # labels-only (no scores): composite None
    res = judge_all(lambda s, u: "good", rubrics, candidate="text")
    assert res["composite"] is None
    assert res["dimensions"]["tone"]["choice"] == "good"

    # garbage answers -> no valid votes
    r = judge_dimension(lambda s, u: "12345 !!!", "tone", rubrics["tone"], "text", k=2)
    assert r["choice"] is None and r["score"] is None

    # voting with disagreement, pair mode
    answers = iter(["good", "bad", "good"])
    r2 = judge_dimension(lambda s, u: next(answers), "tone", rubrics["tone"], "new", original="old", k=3)
    assert r2["choice"] == "good" and r2["agreement"]["winner_share"] == pytest.approx(0.667, abs=0.001)

    assert vote_metrics([], 4)["winner"] is None
    assert vote_metrics(["a"], 1)["a_k"] == 1.0

    # flat file without `dimensions` key
    f2 = tmp_path / "flat.yaml"
    f2.write_text("tone:\n  choices: {a: x}\n")
    assert "tone" in load_rubrics_file(f2)


def test_core_rubrics_wrappers(tmp_path, stub):
    from revoice.core.rubrics import judge_all, load_rubrics, write_template

    rubrics = load_rubrics(tmp_path)  # no file -> template
    assert "voice_fidelity" in rubrics
    assert write_template(tmp_path) is True
    assert write_template(tmp_path) is False
    rubrics2 = load_rubrics(tmp_path)  # file present now
    assert "integrity" in rubrics2
    lines = []
    res = judge_all(stub, rubrics2, "orig text", "new text", k=1,
                    progress=lambda d, m: lines.append(d))
    assert res["rejected_by"] == [] and lines


def test_voicepack_create_drops_protective_gitignore(tmp_path):
    from revoice.core.voicepack import VoicePack

    pack = VoicePack.create(tmp_path / "data", "me")
    gi = pack.root / ".gitignore"
    assert gi.is_file() and "*" in gi.read_text()
    # never overwrites a user's own version
    gi.write_text("custom")
    VoicePack.create(tmp_path / "data", "me")
    assert gi.read_text() == "custom"


# ---------- span-level calibration ----------

def test_bucket_for_and_span_floor_fallbacks(learned_pack):
    """A span longer than any measured band falls back to the nearest SHORTER band.

    Scores climb with length, so a shorter bucket's band is a lower bar. Erring that
    way keeps the author's own text untouched, which is the expensive mistake to avoid.
    """
    from revoice.core.metrics import SPAN_BUCKETS, _bucket_for, load_calibration, span_floor

    assert _bucket_for(10) == SPAN_BUCKETS[0][1]      # below the first bucket
    assert _bucket_for(40) == 40
    assert _bucket_for(41) == 80
    assert _bucket_for(10_000) == SPAN_BUCKETS[-1][1]  # above the last bucket

    calib = load_calibration(learned_pack)
    reg = next(iter(calib))
    bands = sorted(int(k) for k in calib[reg]["spans"])
    assert bands, "the fixture pack must produce span bands"

    # a very long span reuses the largest band we actually measured
    longest = span_floor(calib, reg, 100_000)
    assert longest == pytest.approx(calib[reg]["spans"][str(bands[-1])]["mean"]
                                    - calib[reg]["spans"][str(bands[-1])]["std"])
    # unknown register, and a register with no spans at all -> decline to judge
    assert span_floor(calib, "nonexistent", 100) is None
    assert span_floor({reg: {"spans": {}}}, reg, 100) is None
    # a span shorter than every band still gets the smallest one
    assert span_floor({reg: {"spans": {"320": {"mean": 50.0, "std": 5.0}}}}, reg, 30) \
        == pytest.approx(45.0)


def test_calibration_bands_rise_with_span_length(learned_pack):
    """The whole point: the score is length-dependent, so the band must be too."""
    from revoice.core.metrics import load_calibration

    calib = load_calibration(learned_pack)
    reg = next(iter(calib))
    spans = calib[reg]["spans"]
    means = [spans[k]["mean"] for k in sorted(spans, key=int)]
    assert len(means) >= 2
    assert means[-1] > means[0], "longer spans must score higher than short ones"
    assert calib[reg]["leave_one_out"] is True


def test_baseline_from_texts_handles_missing_histograms():
    """A one-line document has no paragraph histogram; aggregation must not crash."""
    from revoice.core.metrics import baseline_from_texts

    b = baseline_from_texts(["short.", "also short."])
    assert b["doc_count"] == 2
    assert isinstance(b["para_len_hist"], list)


def test_calibration_skips_a_register_with_no_readable_text():
    """An indexed register whose documents have all gone away yields no band, not a crash."""
    from revoice.voicemetric.baseline import baseline_from_texts, calibration_from_texts

    texts = ["Ordinary prose about the apparatus and its behaviour under load. " * 8,
             "Further ordinary prose concerning the same apparatus and its limits. " * 8]
    baselines = {"real": baseline_from_texts(texts), "gone": baseline_from_texts(texts)}
    calib = calibration_from_texts(baselines, {"real": texts, "gone": []})
    assert "real" in calib
    assert "gone" not in calib
